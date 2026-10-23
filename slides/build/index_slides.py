"""Écrit l'index des slides (numéro, titre, source) dans plan/couverture.md.

Usage : uv run python slides/build/index_slides.py
L'index est placé entre les deux marqueurs INDEX_SLIDES de couverture.md.
"""

from pptx import Presentation

from deck import ROOT

COVERAGE = ROOT / "plan" / "couverture.md"
START = "<!-- INDEX_SLIDES_DEBUT -->"
END = "<!-- INDEX_SLIDES_FIN -->"


def slide_title(slide):
    for shape in slide.placeholders:
        if "TITLE" in str(shape.placeholder_format.type) and shape.text_frame.text.strip():
            return shape.text_frame.text.replace("\n", " ").strip()
    for shape in slide.placeholders:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text.replace("\n", " ").strip()
    return ""


def slide_source(slide):
    text = slide.notes_slide.notes_text_frame.text
    for line in text.split("\n"):
        if line.startswith("Source : "):
            return line[len("Source : "):]
    return ""


def index_markdown():
    parts = []
    for day in ("formation",):
        prs = Presentation(str(ROOT / "slides" / f"{day}.pptx"))
        parts.append(f"### slides/{day}.pptx\n\n| N° | Titre | Source |\n|---|---|---|")
        for number, slide in enumerate(prs.slides, start=1):
            title = slide_title(slide).replace("|", "/")
            if slide._element.get("show") == "0":
                title += " (annexe, masquée en diaporama)"
            source = slide_source(slide).replace("|", "/")
            parts.append(f"| {number} | {title} | {source} |")
        parts.append("")
    return "\n".join(parts)


if __name__ == "__main__":
    content = COVERAGE.read_text(encoding="utf-8")
    before, rest = content.split(START)
    _, after = rest.split(END)
    COVERAGE.write_text(f"{before}{START}\n\n{index_markdown()}\n{END}{after}", encoding="utf-8")
    print("Index des slides mis à jour dans plan/couverture.md")
