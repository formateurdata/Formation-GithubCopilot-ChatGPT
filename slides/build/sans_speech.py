"""Copie slides/formation.pptx dans slides/without/formation.pptx, sans le speech des notes du formateur.

Les slides, les animations, la question au groupe et la source sont gardées : seule la partie
« Speech » des notes est retirée. À relancer après chaque build.py.

Usage : uv run python slides/build/sans_speech.py
"""

import sys

from pptx import Presentation

from deck import ROOT

OUTPUT_DIR = ROOT / "slides" / "without"


def sans_speech(notes):
    """Retire le bloc qui commence par « Speech » et s'arrête à la première ligne vide."""
    lines = notes.split("\n")
    if not lines or lines[0] != "Speech":
        return notes
    end = lines.index("") if "" in lines else len(lines)
    return "\n".join(lines[end + 1:])


def copy_day(name):
    prs = Presentation(str(ROOT / "slides" / f"{name}.pptx"))
    for number, slide in enumerate(prs.slides, start=1):
        frame = slide.notes_slide.notes_text_frame
        cleaned = sans_speech(frame.text)
        if "Speech" in cleaned or "(Arrivée de la slide)" in cleaned or "[Clic " in cleaned:
            sys.exit(f"{name}, slide {number} : speech encore présent dans les notes")
        frame.text = cleaned
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / f"{name}.pptx"
    prs.save(str(output))
    # Garder à jour le nom déjà utilisé pour distribuer le support stagiaire.
    if name == "formation":
        prs.save(str(OUTPUT_DIR / "formation-github-chatgpt.pptx"))
    print(f"{output.relative_to(ROOT)} : {len(prs.slides)} slides, notes sans speech")


if __name__ == "__main__":
    for day in sys.argv[1:] or ["formation"]:
        copy_day(day)
